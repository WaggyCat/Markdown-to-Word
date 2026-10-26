#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
mdclean.py —— 检查 Word (.docx) 里残留的 Markdown 标记，并输出清理后的新文件。

典型场景：Markdown 被某个工具/直接粘贴塞进 Word，留下了 `**加粗**`、`## 标题`、
`- 列表`、`| 表格 |`、`[文本](链接)`、`~~删除线~~`、`&nbsp;`、`\\*` 这类纯文本标记。

用法：
    python mdclean.py 输入.docx                  # → 输入.cleaned.docx + 输入.cleaned.cleanreport.md
    python mdclean.py 输入.docx -o 输出.docx --report 报告.md
    python mdclean.py 输入.docx --dry-run        # 只检查、写报告，不产出 docx

可选开关：
    --no-convert-scripts   不把 Unicode 上下标（² ³ ₁ ⁻）转成 Word 上下标
    --remove-cjk-space     顺便删除中英文之间的空格（默认不动）
    --fix-headings         把 `#` 级别写回段落大纲级别（便于自动生成目录）
    --reformat             清理后再套用 md2docx 的标准排版

依赖：python-docx（与 md2docx.py 相同）
    uv run --with python-docx python mdclean.py 输入.docx
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

sys.path.insert(0, str(Path(__file__).resolve().parent))
import md2docx as fmt  # 复用格式函数与上下标映射表

# --------------------------------------------------------------------------- #
# 默认开关（可被命令行覆盖）
# --------------------------------------------------------------------------- #
CONVERT_SCRIPTS = True
REMOVE_CJK_SPACE = False
FIX_HEADINGS = False
REFORMAT = False
KEEP_LIST_MARKERS = False   # True = 不把行首的 `-` / `1.` / `>` 当列表标记删除

SCRIPT_SUP: Set[str] = {c for c in fmt.SUPERSCRIPT_MAP if c not in fmt.SYMBOL_KEEP}
SCRIPT_SUB: Set[str] = set(fmt.SUBSCRIPT_MAP)
ENTITY_MAP = {"nbsp": " ", "amp": "&", "lt": "<", "gt": ">", "quot": "\""}


# --------------------------------------------------------------------------- #
# 规则表：检测与清理共用同一套正则
#   kind:
#     whole    整段都是标记（分隔线 / 表格分隔行）→ 整段删掉
#     prefix   行首前缀（# / - / 1. / >）→ 只删前缀
#     wrap     成对标记 → 删标记、留内容、按需套格式（内容在第 gi 个捕获组）
#     link     链接 / 图片 → 留显示文本
#     strip    逐字符删（HTML 标签）
#     entity   替换成对应字符（按名映射）
#     unescape 只删反斜杠
# --------------------------------------------------------------------------- #
@dataclass
class Rule:
    key: str
    label: str
    regex: re.Pattern
    kind: str
    fmt: Dict[str, str] = field(default_factory=dict)
    gi: int = 1


RULES: Sequence[Rule] = [
    # 整行标记（先处理，免得被前缀/内联规则误伤）
    Rule("tablesep", "`|---|` 表格分隔行",
         re.compile(r"^(?=[^\n]*\|)[ \t]*\|?[ \t:|-]*-[ \t:|-]*\|?[ \t]*$"), "whole"),
    Rule("hr", "`---` / `***` 分隔线",
         re.compile(r"^[ \t]{0,3}([-*_])(?:[ \t]*\1){2,}[ \t]*$"), "whole"),
    # 行首前缀
    Rule("heading", "`#` 标题标记",
         re.compile(r"^[ \t]{0,3}(#{1,6})[ \t]+"), "prefix"),
    Rule("quote", "`>` 引用标记",
         re.compile(r"^[ \t]*>[ \t]?"), "prefix"),
    Rule("ul", "`-` `*` `+` 无序列表标记",
         re.compile(r"^[ \t]*[-*+][ \t]+"), "prefix"),
    Rule("ol", "`1.` 有序列表标记",
         re.compile(r"^[ \t]*\d+[.)][ \t]+"), "prefix"),
    # 表格行
    Rule("tablerow", "`| a | b |` 表格行残留",
         re.compile(r"^[ \t]*\|.*\|[ \t]*$"), "tablerow"),
    # 内联
    Rule("escape", "`\\*` 反斜杠转义",
         re.compile(r"\\([*_\[\]()`#|>~\\-])"), "unescape"),
    Rule("htmltag", "`<br>` 等 HTML 标签",
         re.compile(r"</?(?:br|hr|p|div|span|b|i|u|s|strong|em|sup|sub|code|pre)[ \t]*/?>",
                    re.I), "strip"),
    Rule("entity", "`&nbsp;` 等 HTML 实体",
         re.compile(r"&([a-zA-Z]+|#\d+|#x[0-9a-fA-F]+);"), "entity"),
    Rule("image", "`![alt](src)` 图片标记",
         re.compile(r"!\[([^\]\n]*)\]\([^)\n]*\)"), "link"),
    Rule("link", "`[文本](链接)` 链接标记",
         re.compile(r"\[([^\]\n]*)\]\([^)\n]*\)"), "link"),
    Rule("code", "`` `行内代码` ``",
         re.compile(r"(?<!`)(`{1,3})(?!`)([^\n]+?)(?<!`)\1(?!`)"), "wrap",
         {"code": "1"}, gi=2),
    Rule("bold", "`**加粗**`", re.compile(r"\*\*(?![\s*])(.+?)(?<![\s*])\*\*"),
         "wrap", {"bold": "1"}),
    Rule("bold_u", "`__加粗__`",
         re.compile(r"(?<![\w])__(?![\s_])(.+?)(?<![\s_])__(?![\w])"), "wrap", {"bold": "1"}),
    Rule("strike", "`~~删除线~~`", re.compile(r"~~(?!\s)(.+?)(?<!\s)~~"), "wrap",
         {"strike": "1"}),
    Rule("italic", "`*斜体*`",
         re.compile(r"(?<![*\w])\*(?![\s*])([^*\n]+?)(?<![\s*])\*(?!\*)"), "wrap",
         {"italic": "1"}),
    Rule("italic_u", "`_斜体_`",
         re.compile(r"(?<![\w])_(?![\s_])([^_\n]+?)(?<![\s_])_(?![\w])"), "wrap",
         {"italic": "1"}),
]


# --------------------------------------------------------------------------- #
# 记录结构
# --------------------------------------------------------------------------- #
@dataclass
class Hit:
    where: str
    label: str
    snippet: str
    action: str


@dataclass
class Target:
    para: Paragraph
    where: str


@dataclass
class Plan:
    deleted: Set[int] = field(default_factory=set)
    inserts: Dict[int, str] = field(default_factory=dict)
    formats: List[Tuple[int, int, Dict[str, str]]] = field(default_factory=list)
    consumed: List[Tuple[int, int]] = field(default_factory=list)
    hits: List[Hit] = field(default_factory=list)
    heading_level: Optional[int] = None
    scripts: Dict[int, Tuple[str, str]] = field(default_factory=dict)
    escaped: Set[int] = field(default_factory=set)

    def empty(self) -> bool:
        return not (self.deleted or self.inserts or self.formats or self.scripts
                    or self.heading_level)


def iter_targets(doc) -> List[Target]:
    """按文档顺序收集：正文段落 → 表格单元格段落 → 页眉页脚段落。"""
    targets: List[Target] = []
    for i, p in enumerate(doc.paragraphs, start=1):
        targets.append(Target(p, f"正文 段落#{i}"))
    seen: Set[int] = set()
    for ti, table in enumerate(doc.tables, start=1):
        for ri, row in enumerate(table.rows, start=1):
            for ci, cell in enumerate(row.cells, start=1):
                if id(cell._tc) in seen:        # 合并单元格会重复
                    continue
                seen.add(id(cell._tc))
                for k, p in enumerate(cell.paragraphs, start=1):
                    targets.append(Target(p, f"表格{ti} R{ri}C{ci}#{k}"))
    for si, section in enumerate(doc.sections, start=1):
        for name, hf in (("页眉", section.header), ("页脚", section.footer)):
            try:
                paras = hf.paragraphs
            except Exception:
                continue
            for k, p in enumerate(paras, start=1):
                targets.append(Target(p, f"{name}(第{si}节)#{k}"))
    return targets


def text_with_runs(para: Paragraph) -> Tuple[str, List[Tuple[object, int, int]], bool]:
    """拼接段落里“纯文本 run”的内容，并给出每 run 的字符区间。

    返回 (text, parts, opaque)。opaque=True 表示段落里还有换行/制表/图片/域等
    run（这些 run 被跳过、不参与重建，保证版式不被破坏）。
    """
    chars: List[str] = []
    parts: List[Tuple[object, int, int]] = []
    opaque = False
    offset = 0
    for r in para._p.findall(qn("w:r")):
        children = [c for c in r if c.tag != qn("w:rPr")]
        if not children:
            continue
        if any(c.tag != qn("w:t") for c in children):
            opaque = True
            continue
        txt = "".join((t.text or "") for t in children)
        parts.append((r, offset, offset + len(txt)))
        chars.append(txt)
        offset += len(txt)
    return "".join(chars), parts, opaque


def _snippet(text: str, start: int, end: int, pad: int = 14) -> str:
    a, b = max(0, start - pad), min(len(text), end + pad)
    s = text[a:b].replace("\n", "⏎").replace("|", "\\|").replace("`", "'")
    return ("…" if a > 0 else "") + s + ("…" if b < len(text) else "")


def _overlaps(consumed: List[Tuple[int, int]], s: int, e: int) -> bool:
    return any(not (e <= cs or s >= ce) for cs, ce in consumed)


# --------------------------------------------------------------------------- #
# 分析
# --------------------------------------------------------------------------- #
def analyse(text: str, where: str) -> Plan:
    plan = Plan()
    for rule in RULES:
        for m in rule.regex.finditer(text):
            s, e = m.span()
            if _overlaps(plan.consumed, s, e):
                continue

            if rule.kind == "whole":
                plan.deleted.update(range(s, e))
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e), "已删除整行标记"))

            elif rule.kind == "prefix":
                if text[:s].strip():          # 前缀必须真的在行首
                    continue
                if KEEP_LIST_MARKERS and rule.key in ("ul", "ol", "quote"):
                    continue
                plan.deleted.update(range(s, e))
                plan.consumed.append((s, e))
                if rule.key == "heading":
                    plan.heading_level = len(m.group(1))
                    act = f"已删除『{'#' * plan.heading_level} 』"
                    if FIX_HEADINGS:
                        act += f"，并写入大纲级别 {plan.heading_level}"
                else:
                    act = "已删除行首标记（如需项目符号请套用 Word 列表样式）"
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e), act))

            elif rule.kind == "tablerow":
                # 去掉竖线，并把因此产生的多余空格收成单个空格、去掉首尾空白
                kept: List[int] = []
                prev_space = True          # 行首空白当作“前一个字符是空格”
                for i in range(s, e):
                    ch = text[i]
                    if ch == "|":
                        plan.deleted.add(i)
                        continue
                    if ch in " \t\u00a0\u3000":
                        if prev_space:
                            plan.deleted.add(i)
                        else:
                            prev_space = True
                            kept.append(i)
                    else:
                        prev_space = False
                        kept.append(i)
                while kept and text[kept[-1]] in " \t\u00a0\u3000":   # 去尾部空白
                    plan.deleted.add(kept.pop())
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e),
                                     "已去掉竖线、收拢单元格之间的空格（表格结构需人工重排）"))

            elif rule.kind in ("link", "image"):
                ks, ke = m.span(1)
                plan.deleted.update(range(s, ks))
                plan.deleted.update(range(ke, e))
                plan.consumed.append((s, e))
                note = "已保留显示文本、删除链接/图片标记"
                if rule.kind == "link" and ke - ks == 0:
                    note = "链接文本为空，已删除整个标记"
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e), note))

            elif rule.kind == "strip":
                plan.deleted.update(range(s, e))
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e), "已删除标签"))

            elif rule.kind == "entity":
                name = m.group(1).lower()
                if name.startswith("#"):
                    try:
                        rep = chr(int(name[2:], 16) if name.startswith("#x") else int(name[1:]))
                    except (ValueError, OverflowError):
                        rep = " "
                else:
                    rep = ENTITY_MAP.get(name, "")
                plan.deleted.update(range(s, e))
                plan.inserts[s] = rep
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e),
                                     f"已替换为 {rep!r}"))

            elif rule.kind == "unescape":
                plan.deleted.add(s)
                plan.escaped.add(s + 1)      # 反斜杠后面那个字符是“字面字符”
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e), "已去掉反斜杠"))

            elif rule.kind == "wrap":
                cs, ce = m.span(rule.gi)
                if cs < 0:
                    continue
                plan.deleted.update(range(s, cs))
                plan.deleted.update(range(ce, e))
                plan.formats.append((cs, ce, dict(rule.fmt)))
                plan.consumed.append((s, e))
                plan.hits.append(Hit(where, rule.label, _snippet(text, s, e),
                                     "已去掉标记并套用格式"))

    if CONVERT_SCRIPTS:
        for i, ch in enumerate(text):
            if i in plan.deleted:
                continue
            if ch in SCRIPT_SUP:
                plan.scripts[i] = ("superscript", fmt.SUPERSCRIPT_MAP[ch])
            elif ch in SCRIPT_SUB:
                plan.scripts[i] = ("subscript", fmt.SUBSCRIPT_MAP[ch])
        if plan.scripts:
            first = min(plan.scripts)
            plan.hits.append(Hit(where, "Unicode 上/下标", _snippet(text, first, first + 2),
                                 f"已转换 {len(plan.scripts)} 个字符为 Word 上下标"))

    if REMOVE_CJK_SPACE:
        dropped = 0
        i, n = 0, len(text)
        while i < n:
            if text[i] in fmt.SPACE_CHARS and i not in plan.deleted:
                j = i
                while j < n and text[j] in fmt.SPACE_CHARS:
                    j += 1
                prev = text[i - 1] if i > 0 else ""
                nxt = text[j] if j < n else ""
                if prev and nxt and ((fmt._is_cjk(prev) and fmt._is_latin(nxt))
                                     or (fmt._is_latin(prev) and fmt._is_cjk(nxt))):
                    plan.deleted.update(range(i, j))
                    dropped += 1
                i = j
            else:
                i += 1
        if dropped:
            plan.hits.append(Hit(where, "中英文之间的空格", f"（{dropped} 处）",
                                 f"已删除 {dropped} 处空格"))
    return plan


def detect_only(text: str, protect: Optional[Set[int]] = None,
                skip_prefix: bool = False) -> List[Tuple[str, str]]:
    """只做检测（给复核和含图片/换行等复杂段落用），返回 [(规则名, 片段)]。

    protect：这些下标上的字符是“字面量”（如 \\* 转义而来），不要当标记。
    skip_prefix：跳过行首前缀规则（编号如“1. 引言”本来就是正文内容）。
    """
    protect = protect or set()
    found: List[Tuple[str, str]] = []
    used: List[Tuple[int, int]] = []
    for rule in RULES:
        if skip_prefix and rule.kind == "prefix":
            continue
        for m in rule.regex.finditer(text):
            s, e = m.span()
            if _overlaps(used, s, e):
                continue
            if rule.kind == "prefix" and text[:s].strip():
                continue
            if protect and any(k in protect for k in range(s, e)):
                continue
            used.append((s, e))
            found.append((rule.label, _snippet(text, s, e)))
    return found


# --------------------------------------------------------------------------- #
# 重建段落
# --------------------------------------------------------------------------- #
def _new_run(src_r, text: str, flags: Dict[str, str], parent) -> object:
    r = OxmlElement("w:r")
    src_rpr = src_r.find(qn("w:rPr"))
    if src_rpr is not None:
        r.append(copy.deepcopy(src_rpr))     # 原字体/字号/颜色等全部继承
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    run = Run(r, parent)
    if flags.get("bold"):
        run.bold = True
    if flags.get("italic"):
        run.italic = True
    if flags.get("strike"):
        run.font.strike = True
    va = flags.get("vertalign")
    if va == "superscript":
        run.font.superscript = True
    elif va == "subscript":
        run.font.subscript = True
    if flags.get("code"):
        rPr = run._element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.insert(0, rf)
        for attr in ("ascii", "hAnsi", "cs"):
            rf.set(qn("w:" + attr), "Consolas")
        rf.set(qn("w:eastAsia"), fmt.CN_FONT)
    return r


def rebuild(para: Paragraph, text: str, parts, plan: Plan) -> Tuple[int, List[int]]:
    """按 plan 重写段落里的文本 run（原位替换，保留图片/换行等其它元素）。

    返回 (新建 run 数, 映射表)；映射表第 k 项是“新文本第 k 个字符对应的原文本下标”，
    用它在复核时认出“本来就是字面量”的字符（如 \\* 转义出来的星号）。
    """
    per_char: Dict[int, Dict[str, str]] = {}
    for s, e, flags in plan.formats:
        for i in range(max(0, s), min(e, len(text))):
            per_char.setdefault(i, {}).update(flags)

    owner = [(s, e, r) for (r, s, e) in parts]

    def src_of(i):
        for s, e, r in owner:
            if s <= i < e:
                return r
        return None

    groups: List[Tuple[object, List[str], Tuple]] = []
    mapping: List[int] = []
    for i in range(len(text)):
        src = src_of(i)
        if src is None:
            continue
        if i in plan.deleted:
            chunk = plan.inserts.get(i)
            if not chunk:
                continue
            flags = dict(per_char.get(i, {}))
            sig = tuple(sorted(flags.items()))
            if groups and groups[-1][0] is src and groups[-1][2] == sig:
                groups[-1][1].append(chunk)
            else:
                groups.append((src, [chunk], sig))
            mapping.extend([i] * len(chunk))
            continue
        flags = dict(per_char.get(i, {}))
        out_ch = text[i]
        if i in plan.scripts:
            va, out_ch = plan.scripts[i]
            flags["vertalign"] = va
        sig = tuple(sorted(flags.items()))
        if groups and groups[-1][0] is src and groups[-1][2] == sig:
            groups[-1][1].append(out_ch)
        else:
            groups.append((src, [out_ch], sig))
        mapping.append(i)

    for src, chunks, sig in groups:
        el = _new_run(src, "".join(chunks), dict(sig), para)
        src.addprevious(el)
    for r, _s, _e in parts:
        if r.getparent() is not None:
            r.getparent().remove(r)
    return len(groups), mapping


def clean_text_within_runs(para: Paragraph, text: str, parts, plan: Plan) -> int:
    """保守模式：只在单个 w:t 内部删字符（含图片/换行的段落用），不套用格式。"""
    changed = 0
    for r, s, e in parts:
        idx = s
        touched = False
        for t in [c for c in r if c.tag == qn("w:t")]:
            buf = []
            for ch in (t.text or ""):
                if idx in plan.deleted:
                    rep = plan.inserts.get(idx)
                    if rep:
                        buf.append(rep)
                else:
                    buf.append(ch)
                idx += 1
                touched = True
            new = "".join(buf)
            if new != (t.text or ""):
                t.text = new
                t.set(qn("xml:space"), "preserve")
                changed += 1
        _ = touched
    return changed


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def clean_document(doc) -> Tuple[List[Hit], List[Hit], int, int]:
    fixed: List[Hit] = []
    skipped: List[Hit] = []
    touched = 0
    made_runs = 0

    for target in iter_targets(doc):
        para = target.para
        p_el = para._p
        if p_el.find(qn("w:ins")) is not None or p_el.find(qn("w:del")) is not None:
            skipped.append(Hit(target.where, "含修订痕迹", "", "已跳过（请先接受/拒绝修订）"))
            continue

        text, parts, opaque = text_with_runs(para)
        if not text:
            continue

        plan = analyse(text, target.where)
        if plan.empty():
            continue

        touched += 1
        if opaque:
            clean_text_within_runs(para, text, parts, plan)
            for h in plan.hits:
                h.where = target.where
                h.action += "（段落含图片/换行等，保守模式：只删文本、未套格式）"
                fixed.append(h)
            after, _, _ = text_with_runs(para)
            for label, snip in detect_only(after):
                skipped.append(Hit(target.where, label + "（残留）", snip,
                                   "未修复：标记跨 run 或有图片/换行/域，请人工确认"))
            continue

        n_runs, mapping = rebuild(para, text, parts, plan)
        made_runs += n_runs
        fixed.extend(plan.hits)

        if FIX_HEADINGS and plan.heading_level:
            lvl = fmt.get_or_add_child(p_el.get_or_add_pPr(), "outlineLvl")
            lvl.set(qn("w:val"), str(min(plan.heading_level, 9) - 1))

        # 复核：转义出来的字面字符、原标题里的“1.”编号都不算残留
        protect = {k for k, src_i in enumerate(mapping) if src_i in plan.escaped}
        after, _, _ = text_with_runs(para)
        for label, snip in detect_only(after, protect=protect,
                                       skip_prefix=plan.heading_level is not None):
            skipped.append(Hit(target.where, label + "（残留）", snip, "未修复，需人工确认"))

    return fixed, skipped, touched, made_runs


def reformat_document(doc) -> None:
    """套用 md2docx 的标准排版（宋体/新罗马、两端对齐、首行缩进 2 字符、1.5 倍行距）。"""
    for target in iter_targets(doc):
        p = target.para
        if not p.text.strip():
            continue
        pPr = p._p.pPr
        lvl = None
        if pPr is not None and pPr.find(qn("w:outlineLvl")) is not None:
            lvl = int(pPr.find(qn("w:outlineLvl")).get(qn("w:val"))) + 1
        if lvl:
            spec = fmt.HEADING_LAYOUT.get(min(max(lvl, 1), 6),
                                          {"bold": True, "align": "justify", "indent": 2})
            fmt.apply_para_format(p, first_line=spec["indent"] > 0,
                                  align=fmt.ALIGN_MAP.get(spec["align"]))
            for r in p.runs:
                fmt.apply_font(r, size_pt=fmt.HEADING_SIZE)
                if spec["bold"]:
                    r.bold = True
        else:
            fmt.apply_para_format(p, first_line=True)
            for r in p.runs:
                fmt.apply_font(r, size_pt=fmt.FONT_SIZE)


def write_report(path: Path, src: Path, out: Optional[Path], fixed: List[Hit],
                 skipped: List[Hit], touched: int, made_runs: int, dry: bool) -> None:
    counts: Dict[str, int] = {}
    for h in fixed:
        counts[h.label] = counts.get(h.label, 0) + 1

    L: List[str] = ["# Word 文档 Markdown 残留检查报告\n"]
    L.append(f"- 源文件：`{src}`")
    L.append(f"- 输出文件：{'（仅检查，未输出）' if dry or out is None else f'`{out}`'}")
    L.append(f"- 检查时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    L.append(f"- 涉及段落：{touched} 个；重建文本片段：{made_runs} 个")
    L.append(f"- 命中：{len(fixed)} 处；未处理：{len(skipped)} 处\n")

    if counts:
        L.append("## 命中汇总\n")
        L.append("| 规则 | 次数 |")
        L.append("|---|---|")
        for label, n in sorted(counts.items(), key=lambda kv: -kv[1]):
            L.append(f"| {label} | {n} |")
        L.append("")

    if fixed:
        L.append("## 明细\n")
        L.append("| # | 位置 | 规则 | 原文片段 | 处理 |")
        L.append("|---|---|---|---|---|")
        for i, h in enumerate(fixed, start=1):
            L.append(f"| {i} | {h.where} | {h.label} | `{h.snippet}` | {h.action} |")
        L.append("")

    if skipped:
        L.append("## 未处理 / 需人工确认\n")
        L.append("| # | 位置 | 类型 | 片段 | 说明 |")
        L.append("|---|---|---|---|---|")
        for i, h in enumerate(skipped, start=1):
            L.append(f"| {i} | {h.where} | {h.label} | `{h.snippet}` | {h.action} |")
        L.append("")

    L.append("## 说明\n")
    L.append("- 只处理被识别出的标记，其余文字、字体、图片、表格结构不动；源文件不被修改。")
    L.append("- 列表标记（`-`、`1.`、`>`）删除后不会自动变成 Word 列表，需要项目符号请套用列表样式。")
    L.append("- `| a | b |` 这类表格行只去掉竖线，不会重建成 Word 表格。")
    L.append("- 含图片、换行符、域的段落走保守模式：只删文本标记，不套用加粗/斜体等格式。")
    L.append("- 行首的 `-`、`1.`、`>` 会被当成列表标记删除；如果那是正文编号（如“1. 引言”），"
             "请用 `--keep-list-markers` 重跑。")
    L.append("- 建议同一份文件只跑一次：`\\*` 转成字面量 `*` 后再跑，会被当成斜体标记。")
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


def main(argv: List[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(
        description="检查 docx 里残留的 Markdown 标记，输出清理后的新 docx 与检查报告")
    ap.add_argument("input", help="输入的 .docx")
    ap.add_argument("-o", "--output", default=None, help="输出 .docx（默认 输入.cleaned.docx）")
    ap.add_argument("--report", default=None, help="报告 .md 路径")
    ap.add_argument("--dry-run", "--check-only", dest="dry_run", action="store_true",
                    help="只检查并写报告，不输出 docx")
    ap.add_argument("--no-convert-scripts", dest="convert_scripts", action="store_false",
                    default=None, help="不转换 Unicode 上下标")
    ap.add_argument("--remove-cjk-space", dest="remove_cjk_space", action="store_true",
                    default=None, help="顺便删除中英文之间的空格")
    ap.add_argument("--fix-headings", dest="fix_headings", action="store_true",
                    default=None, help="把 # 级别写回段落大纲级别")
    ap.add_argument("--keep-list-markers", dest="keep_list_markers", action="store_true",
                    default=None, help="不删除行首的 - / 1. / > （当它们其实是正文编号时用）")
    ap.add_argument("--reformat", dest="reformat", action="store_true", default=None,
                    help="清理后套用标准排版（宋体/新罗马、两端对齐、首行缩进 2 字符、1.5 倍行距）")
    args = ap.parse_args(argv)

    global CONVERT_SCRIPTS, REMOVE_CJK_SPACE, FIX_HEADINGS, REFORMAT, KEEP_LIST_MARKERS
    if args.convert_scripts is not None:
        CONVERT_SCRIPTS = args.convert_scripts
    if args.remove_cjk_space is not None:
        REMOVE_CJK_SPACE = args.remove_cjk_space
    if args.fix_headings is not None:
        FIX_HEADINGS = args.fix_headings
    if args.reformat is not None:
        REFORMAT = args.reformat
    if args.keep_list_markers is not None:
        KEEP_LIST_MARKERS = args.keep_list_markers

    src = Path(args.input).expanduser()
    if not src.is_file():
        print(f"错误：找不到文件 {src}", file=sys.stderr)
        return 2
    if src.suffix.lower() != ".docx":
        print("错误：只支持 .docx（.doc 请先另存为 .docx）", file=sys.stderr)
        return 2

    out = Path(args.output).expanduser() if args.output else src.with_suffix(".cleaned.docx")
    if out.resolve() == src.resolve():
        print("错误：输出不能覆盖源文件，请换 -o 路径", file=sys.stderr)
        return 2
    report = Path(args.report).expanduser() if args.report else out.with_suffix(".cleanreport.md")

    doc = Document(str(src))
    fixed, skipped, touched, made_runs = clean_document(doc)
    if REFORMAT:
        reformat_document(doc)

    if not args.dry_run:
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
    write_report(report, src, out, fixed, skipped, touched, made_runs, args.dry_run)

    print(f"[mdclean] 命中 {len(fixed)} 处，未处理 {len(skipped)} 处，涉及段落 {touched} 个")
    if not args.dry_run:
        print(f"[OK] 已输出：{out}")
    print(f"[报告] {report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
