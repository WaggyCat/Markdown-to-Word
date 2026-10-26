#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx.py —— 将 Markdown 文件转换为 Word (.docx)，并套用固定文本格式：

    中文字体：宋体 (SimSun)
    西文字体：Times New Roman
    正文字号：小四 (12 pt)
    标题字号：四号 (14 pt)，多级标题分级别
    对齐：全部两端对齐（代码块除外）
    段前 / 段后：0（正文与标题都是）
    首行缩进：2 字符（w:firstLineChars="200"，随字号自动换算）
    行距：1.5 倍
    Unicode 上下标（² ³ ₁ ⁻ ⁿ …）→ Word 真正的上标 / 下标格式
    中英文之间的空格自动删除

多级标题格式（HEADING_LAYOUT，可自行修改）：
    一级：四号加粗、居中（--heading-justify 改为两端对齐）
    二级：四号加粗、两端对齐、顶格
    三〜六级：四号加粗、两端对齐、首行缩进 2 字符

用法：
    python md2docx.py 输入.md                 # 输出 输入.docx
    python md2docx.py 输入.md -o 输出.docx
    python md2docx.py a.md b.md               # 批量
    可用 --cn-font/--en-font/--size/--heading-size/--align/--line-spacing
        /--first-line-indent/--space-before/--space-after/--heading-justify
        /--no-heading-bold/--indent-code/--indent-table 临时覆盖默认格式

依赖：python-docx
    uv run --with python-docx python md2docx.py 输入.md
    pip install python-docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Iterable, List, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

# --------------------------------------------------------------------------- #
# 固定格式参数（想改格式只改这里）
# --------------------------------------------------------------------------- #
CN_FONT = "宋体"                # 中文
EN_FONT = "Times New Roman"     # 西文（新罗马）
FONT_SIZE = 12                  # 小四 = 12 pt
FIRST_LINE_CHARS = 200          # 首行缩进 2 字符（Word 中 1 字符 = 100）
LINE_SPACING = 1.5              # 1.5 倍行距
SPACE_BEFORE = 0                # 段前 0
SPACE_AFTER = 0                 # 段后 0

HEADING_SIZE = 14               # 标题字号：四号 = 14 pt（各级标题统一）
HEADING_BOLD: bool | None = None  # None = 按级别设定；True/False = 强制全部加粗/不加粗
HEADING_CENTER_L1 = True        # 一级标题居中（False = 与正文一样两端对齐）
CODE_NO_INDENT = True           # 代码块不首行缩进
TABLE_NO_INDENT = True          # 表格内文字不首行缩进
REMOVE_CJK_LATIN_SPACE = True   # 删除中文与英文/数字之间的空格（如“中文 English” → “中文English”）

BODY_ALIGN = "justify"          # 正文对齐：justify（两端对齐）/ left / center / right

# 多级标题格式表：字号统一用 HEADING_SIZE，段前段后 0、1.5 倍行距由全局参数控制，
# 各级之间靠 加粗 / 对齐 / 首行缩进 区分。要自定义改这张表即可。
HEADING_LAYOUT = {
    1: {"bold": True, "align": "center", "indent": 0},   # 章标题：居中加粗
    2: {"bold": True, "align": "justify", "indent": 0},  # 节标题：两端对齐、顶格
    3: {"bold": True, "align": "justify", "indent": 2},  # 三级：首行缩进 2 字符
    4: {"bold": True, "align": "justify", "indent": 2},
    5: {"bold": True, "align": "justify", "indent": 2},
    6: {"bold": True, "align": "justify", "indent": 2},
}
ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# --------------------------------------------------------------------------- #
# 中英混排空格处理
# --------------------------------------------------------------------------- #
SPACE_CHARS = " \t\u00a0\u3000"   # 半角/不换行/全角空格


def _is_cjk(ch: str) -> bool:
    """汉字、CJK 标点、全角字符都算“中文侧”。"""
    o = ord(ch)
    return (0x3400 <= o <= 0x4DBF        # 扩展 A
            or 0x4E00 <= o <= 0x9FFF     # 基本区
            or 0xF900 <= o <= 0xFAFF     # 兼容汉字
            or 0x3001 <= o <= 0x303F     # CJK 标点（不含全角空格）
            or 0xFF00 <= o <= 0xFFEF)    # 全角字符


def _is_latin(ch: str) -> bool:
    # Unicode 上下标（²、₅、⁻ …）实际显示为拉丁字符，归到“英文侧”
    if ch in SUPERSCRIPT_MAP or ch in SUBSCRIPT_MAP:
        return True
    return ("a" <= ch <= "z") or ("A" <= ch <= "Z") or ("0" <= ch <= "9")


def normalize_cjk_spacing(segs: List["Seg"]) -> List["Seg"]:
    """删除中英文交界处的空格。

    在整段拼接后的文本上判断（而不是逐 run），这样即使空格两侧被
    **粗体** / `代码` / 链接等格式边界分开，也能正确删除。
    """
    if not segs:
        return segs
    full = "".join(s.text for s in segs)
    drop = set()
    i, n = 0, len(full)
    while i < n:
        if full[i] in SPACE_CHARS:
            j = i
            while j < n and full[j] in SPACE_CHARS:
                j += 1
            prev = full[i - 1] if i > 0 else ""
            nxt = full[j] if j < n else ""
            if prev and nxt and ((_is_cjk(prev) and _is_latin(nxt))
                                 or (_is_latin(prev) and _is_cjk(nxt))):
                drop.update(range(i, j))
            i = j
        else:
            i += 1
    if not drop:
        return segs
    out, k = [], 0
    for s in segs:
        buf = []
        for ch in s.text:
            if k not in drop:
                buf.append(ch)
            k += 1
        out.append(Seg("".join(buf), s.bold, s.italic, s.strike, s.code, s.link))
    return [s for s in out if s.text]

# --------------------------------------------------------------------------- #
# Unicode 上标 / 下标 字符表
# --------------------------------------------------------------------------- #
SUPERSCRIPT_MAP = {
    "⁰": "0", "¹": "1", "²": "2", "³": "3", "⁴": "4",
    "⁵": "5", "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9",
    "⁺": "+", "⁻": "-", "⁼": "=", "⁽": "(", "⁾": ")",
    "ⁿ": "n", "ⁱ": "i", "ᵃ": "a", "ᵇ": "b", "ᶜ": "c", "ᵈ": "d", "ᵉ": "e",
    "ᶠ": "f", "ᵍ": "g", "ʰ": "h", "ʲ": "j", "ᵏ": "k", "ˡ": "l", "ᵐ": "m",
    "ᵒ": "o", "ᵖ": "p", "ʳ": "r", "ˢ": "s", "ᵗ": "t", "ᵘ": "u", "ᵛ": "v",
    "ʷ": "w", "ˣ": "x", "ʸ": "y", "ᶻ": "z", "ᴬ": "A", "ᴮ": "B", "ᴰ": "D",
    "ᴱ": "E", "ᴳ": "G", "ᴴ": "H", "ᴵ": "I", "ᴶ": "J", "ᴷ": "K", "ᴸ": "L",
    "ᴹ": "M", "ᴺ": "N", "ᴼ": "O", "ᴾ": "P", "ᴿ": "R", "ᵀ": "T", "ᵁ": "U",
    "ⱽ": "V", "ᵂ": "W", "°": "°",  # 度符号本身不变，仅占位
}
SUBSCRIPT_MAP = {
    "₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4",
    "₅": "5", "₆": "6", "₇": "7", "₈": "8", "₉": "9",
    "₊": "+", "₋": "-", "₌": "=", "₍": "(", "₎": ")",
    "ₐ": "a", "ₑ": "e", "ₒ": "o", "ₓ": "x", "ₔ": "e", "ₕ": "h",
    "ₖ": "k", "ₗ": "l", "ₘ": "m", "ₙ": "n", "ₚ": "p", "ₛ": "s",
    "ₜ": "t", "ᵢ": "i", "ᵣ": "r", "ᵤ": "u", "ᵥ": "v", "ⱼ": "j",
}
# 度符号 ° 保持原样（它既是上标又被当作普通符号使用）
SYMBOL_KEEP = {"°"}


def split_script_runs(text: str) -> List[Tuple[str, str]]:
    """把一段文本拆成 [(子串, 'normal'|'sup'|'sub'), ...]，保证顺序连续。"""
    out: List[Tuple[str, str]] = []
    buf: List[str] = []
    cur = "normal"

    def flush():
        nonlocal buf, cur
        if buf:
            out.append(("".join(buf), cur))
            buf = []

    for ch in text:
        if ch in SYMBOL_KEEP:
            kind = "normal"
            rep = ch
        elif ch in SUPERSCRIPT_MAP:
            kind, rep = "sup", SUPERSCRIPT_MAP[ch]
        elif ch in SUBSCRIPT_MAP:
            kind, rep = "sub", SUBSCRIPT_MAP[ch]
        else:
            kind, rep = "normal", ch
        if kind != cur:
            flush()
            cur = kind
        buf.append(rep)
    flush()
    return out


# --------------------------------------------------------------------------- #
# 富文本（行内 Markdown）解析
# --------------------------------------------------------------------------- #
INLINE_PATTERN = re.compile(
    r"(?P<code>`+)(?P<code_body>.+?)(?P=code)"
    r"|(?P<bold>\*\*)(?P<bold_body>.+?)\*\*"
    r"|(?P<bold2>__)(?P<bold2_body>.+?)__"
    r"|(?P<strike>~~)(?P<strike_body>.+?)~~"
    r"|(?P<ital>\*)(?P<ital_body>[^*]+?)\*"
    r"|(?P<ital2>_)(?P<ital2_body>[^_]+?)_"
    r"|!?\[(?P<link_text>[^\]]*)\]\((?P<link_url>[^)\s]+)[^)]*\)"
    r"|<(?P<autolink>https?://[^>\s]+)>"
)


class Seg:
    """一段带格式的文本。"""

    __slots__ = ("text", "bold", "italic", "strike", "code", "link")

    def __init__(self, text, bold=False, italic=False, strike=False, code=False, link=None):
        self.text, self.bold, self.italic = text, bold, italic
        self.strike, self.code, self.link = strike, code, link


def parse_inline(text: str) -> List[Seg]:
    """把一行 Markdown 行内文本拆成带格式的片段列表。"""
    segs: List[Seg] = []
    pos = 0
    for m in INLINE_PATTERN.finditer(text):
        if m.start() > pos:
            segs.append(Seg(text[pos:m.start()]))
        if m.group("bold_body") is not None:
            segs.append(Seg(m.group("bold_body"), bold=True))
        elif m.group("bold2_body") is not None:
            segs.append(Seg(m.group("bold2_body"), bold=True))
        elif m.group("ital_body") is not None:
            segs.append(Seg(m.group("ital_body"), italic=True))
        elif m.group("ital2_body") is not None:
            segs.append(Seg(m.group("ital2_body"), italic=True))
        elif m.group("strike_body") is not None:
            segs.append(Seg(m.group("strike_body"), strike=True))
        elif m.group("code_body") is not None:
            segs.append(Seg(m.group("code_body"), code=True))
        elif m.group("link_text") is not None:
            segs.append(Seg(m.group("link_text"), link=m.group("link_url")))
        elif m.group("autolink") is not None:
            segs.append(Seg(m.group("autolink"), link=m.group("autolink")))
        pos = m.end()
    if pos < len(text):
        segs.append(Seg(text[pos:]))
    return [s for s in segs if s.text != ""]


# --------------------------------------------------------------------------- #
# 字体 / 段落格式
# --------------------------------------------------------------------------- #
# OOXML 规定的 w:pPr / w:rPr 子元素顺序（顺序错了 Word 会提示“无法读取的内容”）
PPR_SEQ = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]
RPR_SEQ = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
    "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
    "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
    "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
    "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
]


def _ordered_child(parent, tag: str, seq: List[str]):
    """在 parent 中按 OOXML 规定的顺序取 / 建子元素 tag。"""
    name = "w:" + tag
    el = parent.find(qn(name))
    if el is not None:
        return el
    el = parent.makeelement(qn(name), {})
    successors = seq[seq.index(name) + 1:] if name in seq else []
    if successors:
        try:
            parent.insert_element_before(el, *successors)
            return el
        except Exception:
            pass
    parent.append(el)
    return el


def get_or_add_child(pPr, tag: str):
    """按 OOXML 规定的子元素顺序取 / 建 pPr 子元素。"""
    return _ordered_child(pPr, tag, PPR_SEQ)


def _strip_theme_fonts(rFonts):
    """删掉主题字体属性，否则显式字体设置会被主题覆盖。"""
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn("w:" + attr)
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    return rFonts


def apply_font(target, size_pt: float | None = None, cn: str | None = None,
               en: str | None = None, color: RGBColor | None = None):
    """把中西文字体、字号写到 run / style 的 rPr 上。"""
    size_pt = FONT_SIZE if size_pt is None else size_pt
    cn = CN_FONT if cn is None else cn
    en = EN_FONT if en is None else en
    rPr = target._element.get_or_add_rPr()
    rFonts = _strip_theme_fonts(_ordered_child(rPr, "rFonts", RPR_SEQ))
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)
    # sz / szCs 以半磅为单位
    half = str(int(round(size_pt * 2)))
    for tag in ("sz", "szCs"):
        _ordered_child(rPr, tag, RPR_SEQ).set(qn("w:val"), half)
    if color is not None:
        _ordered_child(rPr, "color", RPR_SEQ).set(qn("w:val"), str(color))


# --------------------------------------------------------------------------- #
# 段落格式
# --------------------------------------------------------------------------- #
def apply_para_format(p, first_line: bool = True, align=None, left_chars: int = 0):
    """段前段后 0、1.5 倍行距、首行缩进、对齐方式（默认两端对齐）。"""
    pf = p.paragraph_format
    pf.space_before = Pt(SPACE_BEFORE)
    pf.space_after = Pt(SPACE_AFTER)
    pf.line_spacing = LINE_SPACING
    if align is None:
        align = ALIGN_MAP.get(BODY_ALIGN, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    ind = get_or_add_child(pPr, "ind")
    # 清掉已有的缩进设定
    for attr in ("firstLine", "firstLineChars", "hanging", "hangingChars",
                 "left", "leftChars", "right", "rightChars"):
        key = qn("w:" + attr)
        if key in ind.attrib:
            del ind.attrib[key]
    if first_line:
        ind.set(qn("w:firstLineChars"), str(FIRST_LINE_CHARS))   # 2 字符
        ind.set(qn("w:firstLine"), str(int(FONT_SIZE * 20 * FIRST_LINE_CHARS / 100)))  # 兜底 twips
    if left_chars:
        ind.set(qn("w:leftChars"), str(left_chars * 100))
    if len(ind.attrib) == 0:          # 没设置就删掉空壳
        pPr.remove(ind)
    # 中文/西文自动间距（中文排版更自然）
    for tag in ("autoSpaceDE", "autoSpaceDN"):
        get_or_add_child(pPr, tag).set(qn("w:val"), "1")


# --------------------------------------------------------------------------- #
# 写 run
# --------------------------------------------------------------------------- #
def add_segmented_runs(p, segs: Iterable[Seg], base_bold: bool = False,
                       size: float | None = None):
    """把带格式片段写入段落；顺带把 Unicode 上下标转成 Word 上下标。"""
    segs = list(segs)
    if REMOVE_CJK_LATIN_SPACE:
        segs = normalize_cjk_spacing(segs)
    for seg in segs:
        for chunk, kind in split_script_runs(seg.text):
            if chunk == "":
                continue
            run = p.add_run(chunk)
            apply_font(run, size_pt=size)
            run.bold = True if (seg.bold or base_bold) else None
            run.italic = True if seg.italic else None
            if seg.strike:
                run.font.strike = True
            if seg.link:
                run.font.underline = True
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            if kind == "sup":
                run.font.superscript = True
            elif kind == "sub":
                run.font.subscript = True


# --------------------------------------------------------------------------- #
# 块级解析
# --------------------------------------------------------------------------- #
RE_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
RE_FENCE = re.compile(r"^\s*(```|~~~)\s*([\w+-]*)\s*$")
RE_HR = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
RE_UL = re.compile(r"^(\s*)[-*+]\s+(.*)$")
RE_OL = re.compile(r"^(\s*)(\d+[.)])\s+(.*)$")
RE_QUOTE = re.compile(r"^\s{0,3}>\s?(.*)$")
RE_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")


def split_table_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    # 支持 \| 转义
    cells = re.split(r"(?<!\\)\|", line)
    return [c.replace("\\|", "|").strip() for c in cells]


def iter_blocks(md_lines: List[str]):
    i, n = 0, len(md_lines)
    while i < n:
        line = md_lines[i]
        stripped = line.strip()

        if stripped == "":
            i += 1
            continue

        # 围栏代码块
        m = RE_FENCE.match(line)
        if m:
            fence, _lang = m.group(1), m.group(2)
            i += 1
            buf = []
            while i < n and not md_lines[i].strip().startswith(fence):
                buf.append(md_lines[i])
                i += 1
            i += 1  # 跳过结束围栏
            yield ("code", "\n".join(buf))
            continue

        # 标题
        m = RE_HEADING.match(line)
        if m:
            yield ("heading", (len(m.group(1)), m.group(2)))
            i += 1
            continue

        # 分隔线
        if RE_HR.match(line):
            yield ("hr", "")
            i += 1
            continue

        # 表格：当前行含 | 且下一行是分隔行
        if "|" in line and i + 1 < n and RE_TABLE_SEP.match(md_lines[i + 1]):
            header = split_table_row(line)
            i += 2
            rows = []
            while i < n and md_lines[i].strip() and "|" in md_lines[i]:
                rows.append(split_table_row(md_lines[i]))
                i += 1
            yield ("table", (header, rows))
            continue

        # 引用
        if RE_QUOTE.match(line):
            buf = []
            while i < n and RE_QUOTE.match(md_lines[i]):
                buf.append(RE_QUOTE.match(md_lines[i]).group(1))
                i += 1
            yield ("quote", " ".join(x.strip() for x in buf if x.strip()))
            continue

        # 无序列表
        if RE_UL.match(line):
            items = []
            while i < n and RE_UL.match(md_lines[i]):
                mm = RE_UL.match(md_lines[i])
                items.append((len(mm.group(1)), mm.group(2)))
                i += 1
            yield ("ul", items)
            continue

        # 有序列表
        if RE_OL.match(line):
            items = []
            while i < n and RE_OL.match(md_lines[i]):
                mm = RE_OL.match(md_lines[i])
                items.append((len(mm.group(1)), mm.group(2), mm.group(3)))
                i += 1
            yield ("ol", items)
            continue

        # 普通段落（连续非空行合并为一段，行尾两空格 = 强制换行）
        buf = []
        while i < n:
            cur = md_lines[i]
            if cur.strip() == "":
                break
            if (RE_HEADING.match(cur) or RE_FENCE.match(cur) or RE_HR.match(cur)
                    or RE_QUOTE.match(cur) or RE_UL.match(cur) or RE_OL.match(cur)):
                break
            if "|" in cur and i + 1 < n and RE_TABLE_SEP.match(md_lines[i + 1]):
                break
            buf.append(cur)
            i += 1
        yield ("para", buf)


def build_document(md_text: str) -> Document:
    doc = Document()

    # 1) 文档默认字体（docDefaults）
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault"))
    if rpr_default is not None:
        rPr = rpr_default.find(qn("w:rPr"))
        if rPr is None:
            rPr = rpr_default.makeelement(qn("w:rPr"), {})
            rpr_default.append(rPr)
        rFonts = _strip_theme_fonts(_ordered_child(rPr, "rFonts", RPR_SEQ))
        rFonts.set(qn("w:ascii"), EN_FONT)
        rFonts.set(qn("w:hAnsi"), EN_FONT)
        rFonts.set(qn("w:eastAsia"), CN_FONT)
        rFonts.set(qn("w:cs"), EN_FONT)
        for tag in ("sz", "szCs"):
            _ordered_child(rPr, tag, RPR_SEQ).set(qn("w:val"), str(int(FONT_SIZE * 2)))

    # 1b) 文档默认段落格式（新建段落时也符合要求）
    ppr_default = styles_el.find(qn("w:docDefaults") + "/" + qn("w:pPrDefault"))
    if ppr_default is not None:
        pPr = ppr_default.find(qn("w:pPr"))
        if pPr is None:
            pPr = ppr_default.makeelement(qn("w:pPr"), {})
            ppr_default.append(pPr)
        sp = get_or_add_child(pPr, "spacing")
        sp.set(qn("w:before"), str(int(SPACE_BEFORE * 20)))
        sp.set(qn("w:after"), str(int(SPACE_AFTER * 20)))
        sp.set(qn("w:line"), str(int(LINE_SPACING * 240)))
        sp.set(qn("w:lineRule"), "auto")
        ind = get_or_add_child(pPr, "ind")
        ind.set(qn("w:firstLineChars"), str(FIRST_LINE_CHARS))
        ind.set(qn("w:firstLine"), str(int(FONT_SIZE * 20 * FIRST_LINE_CHARS / 100)))

    # 2) Normal / 各类样式
    normal = doc.styles["Normal"]
    apply_font(normal)
    normal.paragraph_format.space_before = Pt(SPACE_BEFORE)
    normal.paragraph_format.space_after = Pt(SPACE_AFTER)
    normal.paragraph_format.line_spacing = LINE_SPACING
    n_ind = get_or_add_child(normal._element.get_or_add_pPr(), "ind")
    n_ind.set(qn("w:firstLineChars"), str(FIRST_LINE_CHARS))
    n_ind.set(qn("w:firstLine"), str(int(FONT_SIZE * 20 * FIRST_LINE_CHARS / 100)))

    # 标题样式：各级均四号，段前后 0、1.5 倍行距
    for lvl in range(1, 7):
        try:
            st = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        apply_font(st, size_pt=HEADING_SIZE, color=RGBColor(0, 0, 0))
        st.paragraph_format.space_before = Pt(SPACE_BEFORE)
        st.paragraph_format.space_after = Pt(SPACE_AFTER)
        st.paragraph_format.line_spacing = LINE_SPACING

    for name in ("Title", "List Bullet", "List Number", "Quote", "Intense Quote"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        apply_font(st, color=RGBColor(0, 0, 0))
        st.paragraph_format.space_before = Pt(SPACE_BEFORE)
        st.paragraph_format.space_after = Pt(SPACE_AFTER)
        st.paragraph_format.line_spacing = LINE_SPACING

    md_lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    blocks = list(iter_blocks(md_lines))
    for bi, (kind, payload) in enumerate(blocks):
        if kind == "heading":
            level, text = payload
            spec = dict(HEADING_LAYOUT.get(min(max(level, 1), 6),
                                           {"bold": True, "align": "justify", "indent": 2}))
            if level == 1 and not HEADING_CENTER_L1:
                spec["align"] = "justify"
            bold = spec["bold"] if HEADING_BOLD is None else HEADING_BOLD
            p = doc.add_paragraph()
            apply_para_format(p, first_line=spec["indent"] > 0,
                              align=ALIGN_MAP.get(spec["align"], ALIGN_MAP[BODY_ALIGN]))
            add_segmented_runs(p, parse_inline(text), base_bold=bold, size=HEADING_SIZE)
            # 标题层级用“大纲级别”标记，便于自动生成目录
            pPr = p._p.get_or_add_pPr()
            get_or_add_child(pPr, "outlineLvl").set(qn("w:val"), str(level - 1))

        elif kind == "para":
            # 行尾两个空格或反斜杠 = 硬换行；否则同一段落内用空格连接
            HARD = "\x00"
            text = ""
            hard_prev = False
            for k, raw in enumerate(payload):
                hard = raw.endswith("  ") or raw.rstrip().endswith("\\")
                line = raw.strip()
                if line.endswith("\\"):
                    line = line[:-1].strip()
                if k and hard_prev:
                    text += HARD
                elif k:
                    text += " "
                text += line
                hard_prev = hard
            p = doc.add_paragraph()
            apply_para_format(p, first_line=True)
            for k, part in enumerate(text.split(HARD)):
                if k > 0:
                    br_run = p.add_run()
                    apply_font(br_run)
                    br_run.add_break()
                add_segmented_runs(p, parse_inline(part))

        elif kind == "code":
            for ln in payload.split("\n"):
                p = doc.add_paragraph()
                apply_para_format(p, first_line=not CODE_NO_INDENT,
                                  align=WD_ALIGN_PARAGRAPH.LEFT)
                if ln.strip() == "":
                    p.add_run("")
                    continue
                run = p.add_run(ln)
                apply_font(run, cn=CN_FONT, en="Consolas")
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        elif kind == "quote":
            p = doc.add_paragraph()
            apply_para_format(p, first_line=True, left_chars=2)
            add_segmented_runs(p, parse_inline(payload))
            for run in p.runs:
                run.italic = True

        elif kind == "ul":
            for _indent, item in payload:
                p = doc.add_paragraph()
                apply_para_format(p, first_line=True)
                add_segmented_runs(p, [Seg("• ")] + parse_inline(item))

        elif kind == "ol":
            for _indent, num, item in payload:
                p = doc.add_paragraph()
                apply_para_format(p, first_line=True)
                add_segmented_runs(p, [Seg(f"{num} ")] + parse_inline(item))

        elif kind == "hr":
            p = doc.add_paragraph()
            apply_para_format(p, first_line=False)
            pPr = p._p.get_or_add_pPr()
            pbdr = get_or_add_child(pPr, "pBdr")
            bottom = pbdr.makeelement(qn("w:bottom"), {})
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)

        elif kind == "table":
            header, rows = payload
            ncols = max([len(header)] + [len(r) for r in rows]) or 1
            table = doc.add_table(rows=1 + len(rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c in range(ncols):
                cell = table.cell(0, c)
                text = header[c] if c < len(header) else ""
                p = cell.paragraphs[0]
                apply_para_format(p, first_line=not TABLE_NO_INDENT,
                                  align=ALIGN_MAP.get(BODY_ALIGN, WD_ALIGN_PARAGRAPH.JUSTIFY))
                add_segmented_runs(p, parse_inline(text), base_bold=True)
            for r, row in enumerate(rows, start=1):
                for c in range(ncols):
                    text = row[c] if c < len(row) else ""
                    p = table.cell(r, c).paragraphs[0]
                    apply_para_format(p, first_line=not TABLE_NO_INDENT,
                                      align=ALIGN_MAP.get(BODY_ALIGN, WD_ALIGN_PARAGRAPH.JUSTIFY))
                    add_segmented_runs(p, parse_inline(text))
            # 仅当下一块也是表格时补一个空段，避免两表粘连（其他情况不多余空行）
            if bi + 1 < len(blocks) and blocks[bi + 1][0] == "table":
                p = doc.add_paragraph()
                apply_para_format(p, first_line=False)

    # 页码之外的空段落清掉尾部多余
    while doc.paragraphs and not doc.paragraphs[-1].text.strip() \
            and len(doc.paragraphs[-1]._p.findall(qn("w:r"))) == 0 \
            and doc.paragraphs[-1]._p.find(qn("w:pBdr")) is None:
        el = doc.paragraphs[-1]._p
        parent = el.getparent()
        if parent is None:
            break
        # 只删末尾连续的空段（最多保留到倒数第二段）
        if len(doc.paragraphs) <= 1:
            break
        parent.remove(el)

    return doc


def convert(md_path: str | Path, docx_path: str | Path | None = None) -> Path:
    md_path = Path(md_path)
    if not md_path.exists():
        raise FileNotFoundError(md_path)
    raw = md_path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            md_text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        md_text = raw.decode("utf-8", errors="replace")

    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")
    doc = build_document(md_text)
    doc.save(str(docx_path))
    return docx_path


def main(argv: List[str]) -> int:
    # Windows 控制台可能是 GBK，统一成 UTF-8 输出，避免中文/符号报错
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Markdown → Word（宋体/新罗马，正文小四、标题四号且分级，两端对齐，0 段距，首行缩进 2 字符，1.5 倍行距，Unicode 上下标转 Word 上下标）")
    ap.add_argument("input", help="输入的 .md 文件（可一次给多个）", nargs="+")
    ap.add_argument("-o", "--output", help="输出 .docx（仅单个输入时可用）", default=None)
    ap.add_argument("--heading-bold", dest="heading_bold", action="store_true", default=None, help="标题全部加粗（默认按级别表）")
    ap.add_argument("--no-heading-bold", dest="heading_bold", action="store_false", default=None, help="标题全部不加粗")
    ap.add_argument("--heading-center", dest="heading_center", action="store_true", default=None, help="一级标题居中（默认开）")
    ap.add_argument("--heading-justify", dest="heading_center", action="store_false", default=None,
                    help="一级标题也两端对齐（不居中）")
    ap.add_argument("--heading-size", type=float, default=None, help="标题字号（pt，默认 14 = 四号）")
    ap.add_argument("--align", default=None,
                    choices=["justify", "left", "center", "right"],
                    help="正文对齐（默认 justify = 两端对齐）")
    ap.add_argument("--cn-font", default=None, help="中文字体（默认 宋体）")
    ap.add_argument("--en-font", default=None, help="西文字体（默认 Times New Roman）")
    ap.add_argument("--size", type=float, default=None, help="字号（pt，默认 12 = 小四）")
    ap.add_argument("--line-spacing", type=float, default=None, help="行距倍数（默认 1.5）")
    ap.add_argument("--first-line-indent", type=float, default=None,
                    help="首行缩进字符数（默认 2；0 = 不缩进）")
    ap.add_argument("--space-before", type=float, default=None, help="段前（pt，默认 0）")
    ap.add_argument("--space-after", type=float, default=None, help="段后（pt，默认 0）")
    ap.add_argument("--indent-code", dest="indent_code", action="store_true", default=None,
                    help="代码块也首行缩进")
    ap.add_argument("--indent-table", dest="indent_table", action="store_true", default=None,
                    help="表格内文字也首行缩进")
    ap.add_argument("--keep-cjk-spacing", dest="remove_cjk_space", action="store_false",
                    default=None, help="保留中英文之间的空格（默认删除）")
    args = ap.parse_args(argv)

    global CN_FONT, EN_FONT, FONT_SIZE, FIRST_LINE_CHARS, LINE_SPACING
    global SPACE_BEFORE, SPACE_AFTER, CODE_NO_INDENT, TABLE_NO_INDENT
    global REMOVE_CJK_LATIN_SPACE, HEADING_SIZE, HEADING_BOLD, HEADING_CENTER_L1, BODY_ALIGN
    if args.heading_bold is not None:
        HEADING_BOLD = args.heading_bold
    if args.heading_center is not None:
        HEADING_CENTER_L1 = args.heading_center
    if args.heading_size:
        HEADING_SIZE = args.heading_size
    if args.align:
        BODY_ALIGN = args.align
    if args.cn_font:
        CN_FONT = args.cn_font
    if args.en_font:
        EN_FONT = args.en_font
    if args.size:
        FONT_SIZE = args.size
    if args.line_spacing:
        LINE_SPACING = args.line_spacing
    if args.first_line_indent is not None:
        FIRST_LINE_CHARS = int(round(args.first_line_indent * 100))
    if args.space_before is not None:
        SPACE_BEFORE = args.space_before
    if args.space_after is not None:
        SPACE_AFTER = args.space_after
    if args.indent_code is not None:
        CODE_NO_INDENT = not args.indent_code
    if args.indent_table is not None:
        TABLE_NO_INDENT = not args.indent_table
    if args.remove_cjk_space is not None:
        REMOVE_CJK_LATIN_SPACE = args.remove_cjk_space

    if args.output and len(args.input) > 1:
        print("错误：-o 只能与单个输入文件一起使用", file=sys.stderr)
        return 2
    for src in args.input:
        out = convert(src, args.output)
        print(f"[OK] {src}  ->  {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
